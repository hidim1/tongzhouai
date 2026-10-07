import sys,json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
obj=json.load(open(sys.argv[1]));r=obj['result'];doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=doc.styles[name];s.font.name='Arial Unicode MS';s.font.color.rgb=RGBColor(0,0,0);s.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
for style in doc.styles:
 if style.type==1:
  pr=style.element.find(qn('w:pPr'))
  if pr is not None:
   for b in list(pr.findall(qn('w:pBdr'))):pr.remove(b)
  if style.element.rPr is not None:
   fonts=style.element.rPr.rFonts
   if fonts is not None:
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+key),None)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(26);doc.styles['Heading 1'].font.size=Pt(16)
sec.header.paragraphs[0].text='同舟纵横（厦门）流体技术有限公司'
sec.footer.paragraphs[0].text='同州 AI 生成技术标初稿  需工程师审核' + ('  演示数据' if obj['mode']=='demo' else '')
doc.add_paragraph('技术标初稿',style='Title');doc.add_paragraph(obj['project'])
doc.add_paragraph('演示模式输出，不作为正式报价或技术承诺。' if obj['mode']=='demo' else '根据现有项目资料整理，未经确认的参数和承诺保留待确认状态。')
doc.add_paragraph('生成依据 '+r['source']);doc.add_paragraph(r['summary'])
for s in r['sections']:
 doc.add_heading(s['title'],level=1)
 for p in s['content'].split('\n'):
  if p.strip():doc.add_paragraph(p.strip())
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v if v is not None else '')
 for ri,row in enumerate(t.rows):
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17384A' if ri==0 else ('F3F6F8' if ri%2==0 else 'FFFFFF'));pr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
   pr.append(borders)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
    for run in p.runs:
     run.font.size=Pt(9)
     if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
 return t
doc.add_heading('附表 用户需求说明书响应表',level=1)
table(['编号','需求内容及来源','供方响应'],[[x['id'],x['content']+'\n来源：'+x['source'],x['status']+'\n'+(x.get('response') or '需工程师确认具体响应')] for x in r['requirements']],[1.9,10.5,4.6])
doc.add_heading('附表 设备配置草稿',level=1)
table(['名称','规格型号','数量','状态'],[[x['name'],x['spec'],x['quantity'],x['status']] for x in r['materials']],[4,8,2,3])
doc.save(sys.argv[2]);print(sys.argv[2])
